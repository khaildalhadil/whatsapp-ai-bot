"""
The pitch bookings table, kept in an Excel file (bookings.xlsx).

Each row is one booked hour:  التاريخ | اليوم | الوقت | الاسم | رقم الهاتف
(date, day, time, name, phone). The sheet is in Arabic, right to left.
The AI uses check_day() and book_slot() to answer customers on WhatsApp.

Run "python bookings.py" to (re)create the file with sample bookings.
"""
import os
import threading
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook

FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "bookings.xlsx")
HEADERS = ["التاريخ", "اليوم", "الوقت", "الاسم", "رقم الهاتف"]

# Arabic day names, in Python's order: weekday() is 0 for Monday ... 6 for Sunday.
ARABIC_DAYS = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

OPEN_HOUR = 16   # the pitch opens at 4 PM ...
CLOSE_HOUR = 24  # ... and closes at midnight, so the last booking starts at 23:00

# Oman time (UTC+4, no daylight saving). Azure servers run on UTC, so we set it ourselves.
LOCAL_TIME = timezone(timedelta(hours=4))

# Two people could try to book the same hour at the same moment.
# The lock makes them wait for each other, so one hour is never booked twice.
lock = threading.Lock()


def now():
    """The current date and time in Oman."""
    return datetime.now(LOCAL_TIME).replace(tzinfo=None)


def local_number(phone):
    """Keep the 8-digit Omani number only: "+968 9123 4567" or "٩١٢٣٤٥٦٧" -> "91234567"."""
    # int(c) also turns Arabic digits like "٩" into 9.
    digits = "".join(str(int(c)) for c in str(phone) if c.isdecimal())
    digits = digits.removeprefix("00")  # 00968... -> 968...
    if len(digits) == 11 and digits.startswith("968"):
        digits = digits[3:]
    return digits


def hour_of(cell):
    """Read the hour from a Time cell: works for "20:00" text and for real Excel times."""
    return int(str(cell).split(":")[0])


def read_bookings():
    """Return every booking in the table as a list of dicts."""
    if not os.path.exists(FILE):
        create_sample_file()
    sheet = load_workbook(FILE).active
    rows = sheet.iter_rows(min_row=2, values_only=True)  # min_row=2 skips the header row
    # str(...)[:10] keeps "2026-10-23" even if Excel turned the cell into a real date.
    return [
        {"date": str(row[0])[:10], "hour": hour_of(row[2]), "name": row[3], "phone": row[4]}
        for row in rows
        if row[0]
    ]


def check_day(date):
    """Return the booked and free hours of one date ("YYYY-MM-DD").

    Other customers' names and phone numbers are NOT returned, only the hours.
    """
    day = datetime.strptime(date, "%Y-%m-%d").date()
    booked = sorted(b["hour"] for b in read_bookings() if b["date"] == day.isoformat())
    free = [h for h in range(OPEN_HOUR, CLOSE_HOUR) if h not in booked]
    return {
        "date": day.isoformat(),
        "day": day.strftime("%A"),
        "booked_hours": [f"{h:02d}:00" for h in booked],
        "free_hours": [f"{h:02d}:00" for h in free],
    }


def book_slot(date, hour, name, phone):
    """Book one hour for a customer. Returns {"ok": True, ...} or {"ok": False, "reason": ...}."""
    day = datetime.strptime(date, "%Y-%m-%d").date()

    if not OPEN_HOUR <= hour < CLOSE_HOUR:
        return {"ok": False, "reason": f"The pitch is open from {OPEN_HOUR}:00 to midnight."}
    if datetime(day.year, day.month, day.day, hour) < now():
        return {"ok": False, "reason": "This time is already in the past."}

    with lock:
        # Check again inside the lock: someone may have booked it a second ago.
        if f"{hour:02d}:00" in check_day(date)["booked_hours"]:
            return {"ok": False, "reason": "This hour is already booked."}

        workbook = load_workbook(FILE)
        row = [day.isoformat(), ARABIC_DAYS[day.weekday()], f"{hour:02d}:00", name, local_number(phone)]
        workbook.active.append(row)
        try:
            workbook.save(FILE)
        except PermissionError:
            # On Windows, Excel locks the file while you have it open.
            return {"ok": False, "reason": "The bookings file is open in Excel. Close it and try again."}

    return {"ok": True, "date": day.isoformat(), "day": day.strftime("%A"), "time": f"{hour:02d}:00", "name": name}


def create_sample_file():
    """Create bookings.xlsx with made-up bookings for the next 7 days (1 to 3 per day)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الحجوزات"
    sheet.sheet_view.rightToLeft = True  # Arabic layout: first column on the right
    sheet.append(HEADERS)
    for column, width in zip("ABCDE", [12, 11, 8, 14, 14]):
        sheet.column_dimensions[column].width = width

    names = ["أحمد", "محمد", "سالم", "يوسف", "علي", "سعيد", "حمد"]
    for d in range(7):
        day = now().date() + timedelta(days=d)
        for i, hour in enumerate([17, 20, 22][: d % 3 + 1]):
            fake_phone = f"9{d}{i}12345"  # made-up 8-digit Omani number
            sheet.append([day.isoformat(), ARABIC_DAYS[day.weekday()], f"{hour:02d}:00", names[(d + i) % 7], fake_phone])

    workbook.save(FILE)


if __name__ == "__main__":
    create_sample_file()
    print(f"Created {FILE} with sample bookings:")
    for booking in read_bookings():
        print(booking)
